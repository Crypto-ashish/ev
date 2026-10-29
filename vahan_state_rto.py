#!/usr/bin/env python3
"""
Vahan: Motor Car & Motor Cab registrations by fuel, STATE-wise and RTO-wise,
for financial years FY2004-05 .. FY2024-25.

Built on the request logic of github.com/RevTpark/vahan-scraper (scripts/api.py),
with three fixes that were needed to get correct data:
  * Year Type = Financial Year (values like "2004-2005")
  * table pagination always restarts from row 0 (the server otherwise remembers
    the last page and silently drops rows)
  * Y-axis = Vehicle Class, X-axis = Fuel, then keep only MOTOR CAR / MOTOR CAB

Run it on your own computer (the Vahan site must be reachable):

  pip install httpx beautifulsoup4 openpyxl
  python3 vahan_state_rto.py --level state            # all 36 states  (~756 queries)
  python3 vahan_state_rto.py --level rto              # every RTO      (~31,000 queries, many hours)
  python3 vahan_state_rto.py --level rto --states "Punjab" "Haryana" "Delhi"
  python3 vahan_state_rto.py --level rto --start-fy 2019 --end-fy 2024
  python3 vahan_state_rto.py --excel                  # build the Excel file from what has been scraped

Output (appended as it goes, so an interrupted run resumes where it stopped):
  vahan_state.csv / vahan_rto.csv  -> one row per (state, rto, FY, vehicle class)
                                      with all 36 Vahan fuel columns + TOTAL
  Vahan_State_RTO_Fuel.xlsx        -> grouped EV / Hybrid / Petrol / Diesel / CNG / LPG
"""

import argparse
import csv
import os
import re
import sys
import time

import httpx
from bs4 import BeautifulSoup

URL = "https://vahan.parivahan.gov.in/vahan4dashboard/vahan/view/reportview.xhtml"
FORM_ID = "masterLayout_formlogin"
CLASSES = ("MOTOR CAR", "MOTOR CAB")
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
AJAX_HEADERS = {
    "Faces-Request": "partial/ajax",
    "X-Requested-With": "XMLHttpRequest",
    "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
    "Accept": "application/xml, text/xml, */*; q=0.01",
    "Origin": "https://vahan.parivahan.gov.in",
    "Referer": URL,
}

FUEL_GROUP = {
    "EV": ["ELECTRIC(BOV)", "PURE EV", "FUEL CELL HYDROGEN", "SOLAR"],
    "Hybrid": ["DIESEL/HYBRID", "PETROL/HYBRID", "PETROL/HYBRID/CNG", "PETROL(E20)/HYBRID",
               "PETROL(E20)/HYBRID/CNG", "STRONG HYBRID EV", "PLUG-IN HYBRID EV"],
    "Petrol": ["PETROL", "PETROL(E20)", "ETHANOL(E100)", "FLEX-FUEL(ETHANOL)",
               "PETROL/METHANOL", "METHANOL"],
    "Diesel": ["DIESEL", "BIO-DIESEL(B100)", "FLEX-FUEL(BIO-DIESEL)", "DI-METHYL ETHER"],
    "CNG": ["CNG ONLY", "PETROL/CNG", "PETROL(E20)/CNG", "BIO-CNG/BIO-GAS", "BIO-METHANE",
            "HCNG", "DUAL DIESEL/CNG", "DUAL DIESEL/BIO CNG", "DUAL DIESEL/LNG", "LNG"],
    "LPG": ["LPG ONLY", "PETROL/LPG", "PETROL(E20)/LPG"],
    "Other / NA": ["NOT APPLICABLE", "HYDROGEN(ICE)"],
}
GROUPS = list(FUEL_GROUP)
GROUP_OF = {f: g for g, fs in FUEL_GROUP.items() for f in fs}


# ── Session / AJAX ────────────────────────────────────────────────────────────

class Vahan:
    def __init__(self):
        self.c = httpx.Client(timeout=90, follow_redirects=True, verify=False,
                              headers={"User-Agent": UA})
        self.reset()

    def reset(self):
        html = self.c.get(URL).text
        self.vs = re.search(r'name="javax\.faces\.ViewState"[^>]*value="([^"]+)"', html).group(1)
        soup = BeautifulSoup(html, "html.parser")
        self.state_name = self.display_name = None
        for sel in soup.find_all("select"):
            vals = [o.get("value", "") for o in sel.find_all("option") if o.get("value")]
            if "-1" in vals and sel.get("name") != "selectedRto_input" and all(
                    v.isalpha() and v.isupper() for v in vals[1:6]):
                self.state_name = sel["name"]
                self.states = {o["value"]: o.get_text(strip=True)
                               for o in sel.find_all("option") if o.get("value") not in ("", "-1")}
            elif vals and set(vals) <= {"T", "L", "C", "A"}:
                self.display_name = sel["name"]
        self.refresh_ids = [b["id"] for b in soup.find_all("button")
                            if b.get_text(strip=True).lower() == "refresh" and b.get("id")]
        self.form = {FORM_ID: FORM_ID, "yaxisVar_input": "Vehicle Class", "xaxisVar_input": "Fuel",
                     "selectedRto_input": "-1", "selectedYearType_input": "F",
                     "selectedYear_input": "", self.display_name: "A", self.state_name: "-1"}
        for src in ("yaxisVar", "xaxisVar", "selectedYearType"):
            self.post(src, "change")
        self.refresh_id = None

    def post(self, source, event=None, execute=None, render="@all", extra=None):
        d = {"javax.faces.partial.ajax": "true", "javax.faces.source": source,
             "javax.faces.partial.execute": execute or source,
             "javax.faces.partial.render": render, **self.form, **(extra or {}),
             "javax.faces.ViewState": self.vs}
        if event:
            d["javax.faces.behavior.event"] = event
            d["javax.faces.partial.event"] = event
        else:
            d[source] = source
        for attempt in range(4):
            try:
                r = self.c.post(URL, headers=AJAX_HEADERS, data=d)
                r.raise_for_status()
                break
            except Exception:
                if attempt == 3:
                    raise
                time.sleep(5 * (attempt + 1))
        m = re.search(r'<update id="[^"]*javax\.faces\.ViewState[^"]*"><!\[CDATA\[(.*?)\]\]>', r.text)
        if m:
            self.vs = m.group(1)
        return r.text

    def select_state(self, code):
        self.form[self.state_name] = code
        self.form["selectedRto_input"] = "-1"
        xml = self.post(self.state_name.replace("_input", ""), "change", render="selectedRto yaxisVar")
        rtos = {}
        for cd in re.findall(r"<!\[CDATA\[(.*?)\]\]>", xml, re.S):
            sel = BeautifulSoup(cd, "html.parser").find("select", {"name": "selectedRto_input"})
            if sel:
                rtos = {o["value"]: o.get_text(strip=True) for o in sel.find_all("option")
                        if o.get("value") not in ("", "-1", "0")}
        # state change can reset axes
        self.form["yaxisVar_input"], self.form["xaxisVar_input"] = "Vehicle Class", "Fuel"
        for src in ("yaxisVar", "xaxisVar", "selectedYearType"):
            self.post(src, "change")
        return rtos

    def table(self, fy):
        """Return (fuel_headers, {class: [36 fuel values..., total]}) for current state/RTO."""
        self.form["selectedYear_input"] = fy
        self.post("selectedYear", "change")
        xml = None
        for rid in ([self.refresh_id] if self.refresh_id else []) + self.refresh_ids:
            t = self.post(rid, execute="@all")
            if "groupingTable" in t:
                self.refresh_id, xml = rid, t
                break
        if xml is None:
            return None, {}
        root = re.search(r'<update id="javax.faces.ViewRoot"><!\[CDATA\[(.*?)\]\]></update>', xml, re.S)
        soup = BeautifulSoup(root.group(1) if root else xml, "html.parser")
        g = soup.find(id="groupingTable")
        if g is None:
            return None, {}
        heads = [th.get_text(strip=True) for th in g.select("thead th")]
        # Vahan header has a phantom "Fuel" label and TOTAL before the fuels
        fuels = [h for h in heads[2:] if h not in ("Fuel", "TOTAL")]

        rows = {}
        first = 0
        while first < 2000:
            pt = self.post("groupingTable", "page", execute="groupingTable", render="groupingTable",
                           extra={"groupingTable_pagination": "true",
                                  "groupingTable_first": str(first),
                                  "groupingTable_rows": "25",
                                  "groupingTable_encodeFeature": "true"})
            m = re.search(r'<update id="groupingTable"><!\[CDATA\[(.*?)\]\]></update>', pt, re.S)
            frag = BeautifulSoup("<table><tbody>" + (m.group(1) if m else "") + "</tbody></table>",
                                 "html.parser")
            page = [[td.get_text(strip=True) for td in tr.find_all("td")] for tr in frag.find_all("tr")]
            page = [p for p in page if len(p) > 2]
            for p in page:
                if p[1] in CLASSES:
                    rows[p[1]] = [int(x.replace(",", "") or 0) for x in p[2:]]
            if len(page) < 25:
                break
            first += 25
        for cls, vals in rows.items():
            if len(vals) != len(fuels) + 1 or sum(vals[:-1]) != vals[-1]:
                print(f"      [WARN] {cls} {fy}: columns don't add up, check this row")
        return fuels, rows


# ── Scrape loop ───────────────────────────────────────────────────────────────

def scrape(args):
    out = f"vahan_{args.level}.csv"
    done = set()
    fuels_known = None
    if os.path.exists(out):
        with open(out, newline="", encoding="utf-8") as f:
            r = csv.reader(f)
            fuels_known = next(r)[5:-1]
            for row in r:
                done.add((row[0], row[2], row[3]))  # state, rto, fy
    fys = [f"{y}-{y + 1}" for y in range(args.start_fy, args.end_fy + 1)]

    v = Vahan()
    states = v.states
    if args.states:
        states = {k: l for k, l in states.items()
                  if any(q.lower() in l.lower() for q in args.states)}
    print(f"{len(states)} states, {len(fys)} financial years, level={args.level}")

    f = open(out, "a", newline="", encoding="utf-8")
    w = csv.writer(f)
    n_calls = 0
    for code, slabel in states.items():
        sname = re.sub(r"\(\d+\)$", "", slabel).strip()
        rtos = v.select_state(code)
        targets = [("-1", "ALL RTOs")] if args.level == "state" else list(rtos.items())
        print(f"\n== {sname}: {len(targets)} target(s)")
        for rcode, rlabel in targets:
            v.form["selectedRto_input"] = rcode
            for fy in fys:
                key = (sname, rlabel, fy)
                if key in done:
                    continue
                try:
                    fuels, rows = v.table(fy)
                except Exception as e:
                    print(f"   [ERR] {sname} / {rlabel} / {fy}: {e} — restarting session")
                    time.sleep(10)
                    v = Vahan()
                    v.select_state(code)
                    v.form["selectedRto_input"] = rcode
                    continue
                if fuels is None:
                    print(f"   [skip] {sname} / {rlabel} / {fy}: no table")
                    continue
                if fuels_known is None:
                    fuels_known = fuels
                    w.writerow(["State", "State code", "RTO", "Financial Year", "Vehicle Class"]
                               + fuels_known + ["TOTAL"])
                for cls in CLASSES:
                    vals = rows.get(cls)
                    if vals is None:
                        vals = [0] * (len(fuels) + 1)
                    by = dict(zip(fuels, vals[:-1]))
                    w.writerow([sname, code, rlabel, fy, cls.title()]
                               + [by.get(fu, 0) for fu in fuels_known] + [vals[-1]])
                f.flush()
                done.add(key)
                n_calls += 1
                print(f"   {sname} / {rlabel} / {fy}: car={rows.get('MOTOR CAR', [0])[-1]:,} "
                      f"cab={rows.get('MOTOR CAB', [0])[-1]:,}")
                time.sleep(args.delay)
    f.close()
    print(f"\nDone. {n_calls} new tables saved to {out}")


# ── Excel ─────────────────────────────────────────────────────────────────────

def build_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    wb.remove(wb.active)
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    body = Font(name="Arial", size=10)

    for level in ("state", "rto"):
        path = f"vahan_{level}.csv"
        if not os.path.exists(path):
            continue
        with open(path, newline="", encoding="utf-8") as f:
            r = csv.reader(f)
            head = next(r)
            data = list(r)
        fuels = head[5:-1]
        for cls in ("Motor Car", "Motor Cab"):
            ws = wb.create_sheet(f"{level.upper()}_{cls.replace(' ', '_')}")
            cols = ["State", "RTO", "Financial Year"] + GROUPS + ["Total", "EV %", "Hybrid %"]
            if level == "state":
                cols.remove("RTO")
            ws.append(cols)
            for c in ws[1]:
                c.font, c.fill = hdr_font, hdr_fill
                c.alignment = Alignment(horizontal="center", wrap_text=True)
            for row in data:
                if row[4] != cls:
                    continue
                vals = dict(zip(fuels, map(int, row[5:-1])))
                g = {k: 0 for k in GROUPS}
                for fu, n in vals.items():
                    g[GROUP_OF.get(fu, "Other / NA")] += n
                i = ws.max_row + 1
                lead = [row[0], row[3]] if level == "state" else [row[0], row[2], row[3]]
                ws.append(lead + [g[k] for k in GROUPS])
                first_num = len(lead) + 1
                from openpyxl.utils import get_column_letter as L
                a, b = L(first_num), L(first_num + len(GROUPS) - 1)
                t = L(first_num + len(GROUPS))
                ws.cell(i, first_num + len(GROUPS), f"=SUM({a}{i}:{b}{i})")
                ws.cell(i, first_num + len(GROUPS) + 1, f"=IF({t}{i}=0,0,{a}{i}/{t}{i})")
                ws.cell(i, first_num + len(GROUPS) + 2,
                        f"=IF({t}{i}=0,0,{L(first_num + 1)}{i}/{t}{i})")
                for c in ws[i]:
                    c.font = body
                for j in range(first_num, first_num + len(GROUPS) + 1):
                    ws.cell(i, j).number_format = "#,##0"
                for j in (first_num + len(GROUPS) + 1, first_num + len(GROUPS) + 2):
                    ws.cell(i, j).number_format = "0.0%"
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = 14
            ws.column_dimensions["A"].width = 22
            if level == "rto":
                ws.column_dimensions["B"].width = 34

    m = wb.create_sheet("Fuel_Mapping")
    m.append(["Fuel group", "Vahan fuel types included"])
    for g, fs in FUEL_GROUP.items():
        m.append([g, ", ".join(fs)])
    m.append([])
    m.append(["Source: VAHAN dashboard, Vehicle Class x Fuel, Financial Year, actual values. "
              "Early years are partial: many RTOs joined Vahan 4 only around 2015."])
    m.column_dimensions["A"].width = 16
    m.column_dimensions["B"].width = 110
    wb.save("Vahan_State_RTO_Fuel.xlsx")
    print("Saved Vahan_State_RTO_Fuel.xlsx")


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--level", choices=["state", "rto"], help="state-wise or RTO-wise")
    p.add_argument("--states", nargs="+", help="only these states (partial names)")
    p.add_argument("--start-fy", type=int, default=2004, help="2004 means FY2004-05")
    p.add_argument("--end-fy", type=int, default=2024, help="2024 means FY2024-25")
    p.add_argument("--delay", type=float, default=0.5, help="seconds between tables")
    p.add_argument("--excel", action="store_true", help="build the Excel file from the CSVs")
    a = p.parse_args()
    if a.excel:
        build_excel()
    elif a.level:
        scrape(a)
    else:
        p.error("give --level state|rto, or --excel")


if __name__ == "__main__":
    main()
